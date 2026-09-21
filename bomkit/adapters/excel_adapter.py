import openpyxl
import re
from bomkit.schema import COLUMN_MAPPINGS, STANDARD_HEADERS
from pathlib import Path


class ExcelAdapter:
    def can_handle(self, file_path):
        return Path(file_path).suffix.lower() in [".xlsx", ".xls"]

    def _normalize_header_text(self, text):
        text = str(text).lower().strip()
        text = re.sub(r'[\s_\-]+', ' ', text)
        return text

    def _build_header_aliases(self):
        aliases = {
            self._normalize_header_text(header)
            for header in STANDARD_HEADERS
        }

        for variations in COLUMN_MAPPINGS.values():
            for variation in variations:
                aliases.add(self._normalize_header_text(variation))

        return aliases

    def _score_header_row(self, row, header_aliases):
        if not row:
            return 0.0, 0

        normalized_cells = [
            self._normalize_header_text(cell)
            for cell in row
            if cell is not None
        ]
        non_empty = [cell for cell in normalized_cells if cell]

        if len(non_empty) < 2:
            return 0.0, 0

        header_hits = 0
        partial_hits = 0
        alpha_cells = 0
        numeric_cells = 0

        for cell in non_empty:
            if cell in header_aliases:
                header_hits += 1
            else:
                for alias in header_aliases:
                    if alias and (alias in cell or cell in alias):
                        partial_hits += 1
                        break

            if re.search(r'[a-zA-Z]', cell):
                alpha_cells += 1

            if re.fullmatch(r'[\d\.\-]+', cell):
                numeric_cells += 1

        score = 0.0
        score += header_hits * 2.0
        score += partial_hits * 0.5
        score += alpha_cells * 0.3
        score -= numeric_cells * 0.5
        score += (len(non_empty) / max(len(row), 1)) * 0.2

        return score, header_hits

    def _select_header_row(self, rows):
        if not rows:
            return -1

        header_aliases = self._build_header_aliases()
        best_idx = -1
        best_score = 0.0
        best_hits = 0

        for idx, row in enumerate(rows[:50]):
            score, hits = self._score_header_row(row, header_aliases)

            if score > best_score:
                best_score = score
                best_hits = hits
                best_idx = idx

        if best_idx >= 0 and (best_score >= 2.0 or best_hits >= 2):
            return best_idx

        return -1

    def read(self, file_path):
        wb = openpyxl.load_workbook(file_path, data_only=True)
        header_row = None
        selected_ws = None

        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))[:50]
            header_idx = self._select_header_row(rows)

            if header_idx >= 0:
                header_row = header_idx
                selected_ws = ws
                break

        if selected_ws is None:
            return []

        rows = []
        headers = [
            cell.value
            for cell in selected_ws[header_row + 1]
        ]

        for row in selected_ws.iter_rows(
            min_row=header_row + 2,
            values_only=True
        ):
            row_dict = dict(zip(headers, row))
            rows.append(row_dict)

        return rows